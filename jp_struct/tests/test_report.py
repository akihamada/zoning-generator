#!/usr/bin/env python3
"""
検定表XLSX出力テスト

実際にXLSXを生成し、ファイルが正常に読み取れることを確認する。
実行: python tests/test_report.py (jp_struct/ ディレクトリから)
"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


def test_generate_report():
    """XLSXレポート生成テスト"""
    from jp_struct.sections import get_h_section, get_bh_section, get_box_section, get_material
    from jp_struct.check import MemberInput, check_member
    from jp_struct.report import generate_report

    # 複数の部材を検定
    results = []

    # 梁（H形鋼）
    sec_h = get_h_section("H-400x200x8x13")
    mat_sn = get_material("SN400B")
    inp = MemberInput(
        member_id="B1", section=sec_h, material=mat_sn,
        L=6000, lb=3000, Mx=200e6, Vy=100e3, duration="long",
    )
    results.append(check_member(inp))

    # 梁（BH断面）— NG例
    sec_bh = get_bh_section("BH-500x200x9x16")
    inp = MemberInput(
        member_id="B2", section=sec_bh, material=mat_sn,
        L=8000, lb=4000, Mx=300e6, Vy=120e3, duration="long",
    )
    results.append(check_member(inp))

    # 柱（角形鋼管）
    sec_box = get_box_section("□-300x300x12", "bcr295")
    mat_bcr = get_material("BCR295")
    inp = MemberInput(
        member_id="C1", section=sec_box, material=mat_bcr,
        L=4000, lb=4000, N=-800e3, Mx=50e6, Vy=30e3, duration="long",
    )
    results.append(check_member(inp))

    # 柱（H形鋼）
    sec_h300 = get_h_section("H-300x300x10x15")
    inp = MemberInput(
        member_id="C2", section=sec_h300, material=mat_sn,
        L=4000, lb=4000, N=-500e3, Mx=80e6, Vy=40e3, duration="long",
    )
    results.append(check_member(inp))

    # 短期
    inp = MemberInput(
        member_id="B1-S", section=sec_h, material=mat_sn,
        L=6000, lb=3000, Mx=300e6, Vy=150e3, duration="short",
    )
    results.append(check_member(inp))

    # 一時ファイルに出力
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        filepath = f.name

    try:
        output = generate_report(
            results,
            filepath,
            project_name="テストプロジェクト",
            author="jp_struct テスト",
        )

        # ファイルが存在するか
        assert os.path.exists(output), f"ファイルが見つかりません: {output}"

        # ファイルサイズが妥当か（空でない）
        size = os.path.getsize(output)
        assert size > 1000, f"ファイルサイズが小さすぎます: {size} bytes"

        # openpyxlで読み取れるか
        import openpyxl
        wb = openpyxl.load_workbook(output)
        sheet_names = wb.sheetnames
        assert "部材一覧" in sheet_names, f"シート '部材一覧' がありません: {sheet_names}"
        assert "詳細" in sheet_names, f"シート '詳細' がありません"
        assert "凡例" in sheet_names, f"シート '凡例' がありません"

        # 部材一覧シートのデータ件数
        ws1 = wb["部材一覧"]
        data_rows = ws1.max_row - 5  # ヘッダ5行分を除く
        assert data_rows == len(results), \
            f"データ行数: {data_rows} (期待: {len(results)})"

        # 詳細シートのデータ件数
        ws2 = wb["詳細"]
        data_rows_2 = ws2.max_row - 1  # ヘッダ1行
        assert data_rows_2 == len(results), \
            f"詳細データ行数: {data_rows_2} (期待: {len(results)})"

        wb.close()
        print(f"  OK: XLSX生成成功 ({size} bytes, {len(results)} 部材)")
        print(f"      シート: {sheet_names}")
        print(f"      出力先: {output}")

    finally:
        # テスト後にクリーンアップ
        if os.path.exists(filepath):
            os.remove(filepath)
            print(f"      テンプファイル削除済")


def run_all():
    print("=" * 60)
    print("jp_struct レポート出力テスト")
    print("=" * 60)

    tests = [
        ("XLSX生成", test_generate_report),
    ]

    passed = 0
    failed = 0
    for name, func in tests:
        try:
            print(f"\n[{name}]")
            func()
            passed += 1
        except Exception as e:
            print(f"  NG: {e}")
            import traceback
            traceback.print_exc()
            failed += 1

    print("\n" + "=" * 60)
    print(f"結果: {passed} passed, {failed} failed / {len(tests)} total")
    print("=" * 60)
    return failed == 0


if __name__ == "__main__":
    success = run_all()
    sys.exit(0 if success else 1)
